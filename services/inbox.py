"""Temporary, per-chat Redis collections and bounded document normalization."""
import asyncio
import base64
import io
import json
import re
import zipfile
from pathlib import PurePosixPath

from defusedxml import ElementTree

from services.attachments import Attachment, AttachmentError, UnsupportedAttachment, ensure_size

COLLECTION_TTL = 900
MAX_COLLECTION_ITEMS = 12
MAX_COLLECTION_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 100000


def safe_filename(name):
    name = PurePosixPath(str(name or "attachment").replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f<>:\"|?*]", "_", name).strip(". ")
    return name[:120] or "attachment"


def bounded_text(value):
    if len(value) > MAX_TEXT_CHARS:
        raise AttachmentError("Документ слишком длинный для одного анализа. Раздели его на части.")
    return value


def _normalize(data, filename, mime):
    ensure_size(len(data))
    if not data:
        raise AttachmentError("Пустой файл.")
    filename = safe_filename(filename)
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 20 * 1024 * 1024:
                    raise AttachmentError("DOCX слишком большой после распаковки.")
                root = ElementTree.fromstring(archive.read("word/document.xml"))
                text = "\n".join("".join(p.itertext()) for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"))
        except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
            raise AttachmentError("Повреждённый DOCX.") from exc
        return bounded_text(text), None
    if suffix in {".xlsx", ".xls"}:
        rows = []
        if suffix == ".xlsx":
            import openpyxl
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 20 * 1024 * 1024:
                    raise AttachmentError("Таблица слишком большая после распаковки.")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book.worksheets:
                    rows.append(f"Sheet: {sheet.title}")
                    if sheet.max_column and sheet.max_column > 100:
                        raise AttachmentError("Таблица превышает 100 столбцов.")
                    for i, values in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= 3000:
                            raise AttachmentError("Таблица превышает 3000 строк на лист.")
                        rows.append("\t".join(str(v) if v is not None else "" for v in values))
                        if sum(map(len, rows)) > MAX_TEXT_CHARS:
                            raise AttachmentError("Слишком много данных в таблице.")
            finally:
                book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                for sheet in book.sheets():
                    if sheet.nrows > 3000 or sheet.ncols > 100:
                        raise AttachmentError("Таблица превышает лимит строк/столбцов.")
                    rows.append(f"Sheet: {sheet.name}")
                    rows.extend("\t".join(map(str, sheet.row_values(i))) for i in range(sheet.nrows))
            finally:
                book.release_resources()
        return bounded_text("\n".join(rows)), None
    if suffix == ".pdf" or mime == "application/pdf":
        from pypdf import PdfReader
        if not data.startswith(b"%PDF-"):
            raise AttachmentError("Повреждённый PDF.")
        try:
            pdf = PdfReader(io.BytesIO(data), strict=False)
            if pdf.is_encrypted:
                raise AttachmentError("PDF зашифрован. Нужна копия без пароля.")
            if len(pdf.pages) > 100:
                raise AttachmentError("PDF превышает 100 страниц.")
            text = bounded_text("\n".join(page.extract_text() or "" for page in pdf.pages))
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("Не удалось прочитать PDF.") from exc
        # Keep the original for diagrams and scanned pages; never silently lose
        # visual content by handing only extraction to a text-only provider.
        return text, Attachment(data, filename, "application/pdf", "document")
    if mime.startswith("image/"):
        if mime not in {"image/jpeg", "image/png", "image/webp", "image/gif"}:
            raise UnsupportedAttachment("Поддерживаются JPEG, PNG, WebP и GIF.")
        signatures = {"image/jpeg": data.startswith(b"\xff\xd8\xff"), "image/png": data.startswith(b"\x89PNG\r\n\x1a\n"),
                      "image/webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP", "image/gif": data.startswith((b"GIF87a", b"GIF89a"))}
        if not signatures[mime]:
            raise AttachmentError("Тип изображения не соответствует содержимому.")
        return "", Attachment(data, filename, mime, "image")
    if suffix == ".doc":
        # Legacy binary Word is handled only by a provider advertising native
        # documents. It is never misrepresented as plain text on fallback.
        return "", Attachment(data, filename, "application/msword", "document")
    if mime.startswith("text/") or suffix in {".txt", ".md", ".json", ".csv", ".tsv", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".css", ".html", ".xml", ".sql", ".yaml", ".yml"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise AttachmentError("Текстовый файл должен быть UTF-8.") from exc
        if "\x00" in text:
            raise UnsupportedAttachment("Бинарный файл не является текстом.")
        return bounded_text(text), None
    raise UnsupportedAttachment("Этот тип файла не поддерживается. Используй PDF, DOC/DOCX, изображение, текст или таблицу.")


async def normalize(data, filename, mime):
    try:
        return await asyncio.wait_for(asyncio.to_thread(_normalize, data, filename, mime), timeout=20)
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError("Не удалось прочитать файл безопасно.") from exc


class CollectionStore:
    def __init__(self, redis):
        self.redis = redis

    def key(self, user_id, chat_id):
        return f"jarvis:collection:{user_id}:{chat_id}"

    async def start(self, user_id, chat_id):
        if self.redis is None:
            raise AttachmentError("Сбор файлов временно недоступен: Redis не подключён. Отправляй материалы по одному.")
        key = self.key(user_id, chat_id)
        # An existing collection must never be erased by an accidental /collect.
        await self.redis.set(key, json.dumps({"items": []}), ex=COLLECTION_TTL, nx=True)

    async def active(self, user_id, chat_id):
        return bool(self.redis and await self.redis.get(self.key(user_id, chat_id)))

    async def append(self, user_id, chat_id, text, attachments):
        key = self.key(user_id, chat_id)
        item = {"text": bounded_text(text), "files": [dict(filename=f.filename, mime_type=f.mime_type,
            kind=f.kind, data=base64.b64encode(f.data).decode()) for f in attachments]}
        encoded = json.dumps(item, ensure_ascii=False)
        # Redis Lua serializes concurrent incoming messages and enforces limits
        # atomically. TTL is finite and refreshed only when a new item is added.
        script = """
local raw = redis.call('GET', KEYS[1])
if not raw then return -1 end
local state = cjson.decode(raw)
if #state.items >= tonumber(ARGV[2]) then return -2 end
local count = 0
for _,item in ipairs(state.items) do count = count + string.len(item.text) end
if count + string.len(cjson.decode(ARGV[1]).text) > 90000 then return -3 end
table.insert(state.items, cjson.decode(ARGV[1]))
local data = cjson.encode(state)
if string.len(data) > tonumber(ARGV[3]) then return -3 end
redis.call('SET', KEYS[1], data, 'EX', ARGV[4])
return #state.items
"""
        count = await self.redis.eval(script, keys=[key], args=[encoded, MAX_COLLECTION_ITEMS, MAX_COLLECTION_BYTES, COLLECTION_TTL])
        if count == -1:
            raise AttachmentError("Сбор истёк. Начни /collect заново.")
        if count < 0:
            raise AttachmentError("Достигнут лимит сбора: 12 сообщений, 12 MB. Отправь /done для анализа собранного.")
        return count

    async def finish(self, user_id, chat_id):
        if self.redis is None:
            return None
        raw = await self.redis.get(self.key(user_id, chat_id))
        if not raw:
            return None
        self.snapshot = raw
        state = json.loads(raw)
        texts, files = [], []
        for index, item in enumerate(state["items"]):
            texts.append(f"Material {index + 1}:\n{item['text']}")
            files.extend(Attachment(base64.b64decode(f["data"], validate=True), f["filename"], f["mime_type"], f["kind"]) for f in item["files"])
        return bounded_text("\n\n".join(texts)), files

    async def acknowledge(self, user_id, chat_id):
        # New materials arriving during inference must never be erased.
        return await self.redis.eval("""
if redis.call('GET', KEYS[1]) == ARGV[1] then
  return redis.call('DEL', KEYS[1])
end
return 0
""", keys=[self.key(user_id, chat_id)], args=[self.snapshot])

    async def cancel(self, user_id, chat_id):
        if self.redis is not None:
            await self.redis.delete(self.key(user_id, chat_id))


def starts_collection(text):
    return text.strip().lower() == "/collect" or bool(re.fullmatch(r"сейчас скину несколько (?:сообщений|файлов|сообщений/файлов)[.!]?", text.strip().lower()))


def finishes_collection(text):
    return text.strip().lower() in {"/done", "готово, анализируй", "готово анализируй"}
